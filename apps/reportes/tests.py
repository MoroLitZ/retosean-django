from decimal import Decimal

from django.test import TestCase
from django.urls import reverse

from apps.empresas.models import ContactoEmpresa, DocumentoEmpresa, Empresa
from apps.evaluacion.models import ComentarioEntregable, CriterioRubrica, Entregable, Evaluacion, Rubrica
from apps.participaciones.models import Postulacion
from apps.presupuesto.models import Gasto, Presupuesto
from apps.retos.models import Reto
from apps.usuarios.models import Usuario

from .catalogo import obtener
from .models import ReporteGenerado
from .services import construir_reporte


class GeneracionDeReportesTests(TestCase):
    def setUp(self):
        self.admin = Usuario.objects.create_superuser(
            username="admin-rep", password="x", email="admin@ean.edu.co"
        )
        self.empresa = Usuario.objects.create_user(
            username="empresa-rep", password="x", rol="EMPRESA"
        )
        self.estudiante = Usuario.objects.create_user(
            username="est-rep", password="x", rol="ESTUDIANTE"
        )
        self.profesor = Usuario.objects.create_user(
            username="prof-rep", password="x", rol="PROFESOR"
        )
        self.reto = Reto.objects.create(
            empresa=self.empresa, titulo="Reto reportable", estado="aprobado"
        )
        Postulacion.objects.create(
            reto=self.reto, estudiante=self.estudiante, estado="ACEPTADA", semestre=6,
            motivacion="Quiero aprender", habilidades="Python",
        )
        self.filtros = {"desde": None, "hasta": None}

    def _sembrar_datos_ricos(self):
        """Datos para ejercitar TODAS las columnas nuevas de los 5 reportes."""
        entregable = Entregable.objects.create(
            reto=self.reto, estudiante=self.estudiante, titulo="Avance 1",
            puntaje_maximo=Decimal("5.00"), estado="CALIFICADO", nota=Decimal("4.5"),
            comentario_estudiante="Aqui esta mi avance", comentario_profesor="Buen trabajo",
            archivo="entregables/avance.pdf",
        )
        ComentarioEntregable.objects.create(entregable=entregable, autor=self.profesor, comentario="Recibido.")
        rubrica = Rubrica.objects.create(reto=self.reto, nombre="Rubrica 1", creada_por=self.profesor)
        criterio = CriterioRubrica.objects.create(rubrica=rubrica, nombre="Impacto", puntaje_maximo=Decimal("5.00"))
        Evaluacion.objects.create(
            entregable=entregable, profesor=self.profesor, nota=Decimal("4.5"), rubrica=rubrica,
        )

        presupuesto = Presupuesto.objects.create(reto=self.reto, monto_total=Decimal("1000.00"))
        Gasto.objects.create(
            presupuesto=presupuesto, categoria="MATERIALES", monto=Decimal("200.00"),
            estado=Gasto.COMPROMETIDO,
        )
        Gasto.objects.create(
            presupuesto=presupuesto, categoria="REFRIGERIOS", monto=Decimal("100.00"),
            estado=Gasto.EJECUTADO,
        )

        empresa_perfil = Empresa.objects.create(
            usuario=self.empresa, nit="900123456-1", razon_social="Empresa Reportable",
            sector_industrial="Tecnologia", estado_validacion="VERIFICADA",
        )
        ContactoEmpresa.objects.create(
            empresa=empresa_perfil, nombre="Contacto Uno", cargo="Gerente",
            email="contacto@empresa.co", telefono="3000000000", es_principal=True,
        )
        DocumentoEmpresa.objects.create(empresa=empresa_perfil, tipo_documento="RUT", estado="VERIFICADO")

    def test_genera_pdf_y_registra_el_log(self):
        definicion = obtener("retos")
        registro, contenido, nombre, tipo = construir_reporte(
            self.admin, definicion, definicion.columnas, "PDF", self.filtros
        )
        self.assertTrue(contenido.startswith(b"%PDF"))
        self.assertEqual(tipo, "application/pdf")
        self.assertTrue(nombre.endswith(".pdf"))
        self.assertEqual(registro.usuario, self.admin)
        self.assertEqual(registro.filtros["reporte"], "retos")
        self.assertEqual(registro.filtros["total_registros"], 1)
        self.assertTrue(registro.archivo)

    def test_genera_excel(self):
        definicion = obtener("participacion")
        _, contenido, nombre, tipo = construir_reporte(
            self.admin, definicion, definicion.columnas, "EXCEL", self.filtros
        )
        # Un .xlsx es un zip: empieza por PK.
        self.assertTrue(contenido.startswith(b"PK"))
        self.assertTrue(nombre.endswith(".xlsx"))
        self.assertIn("spreadsheetml", tipo)

    def test_solo_incluye_las_columnas_elegidas(self):
        definicion = obtener("retos")
        columnas = definicion.columnas_por_clave(["titulo", "estado"])
        self.assertEqual([c.clave for c in columnas], ["titulo", "estado"])

    def test_sin_columnas_elegidas_se_incluyen_todas(self):
        definicion = obtener("retos")
        self.assertEqual(
            len(definicion.columnas_por_clave([])), len(definicion.columnas)
        )

    def test_todos_los_reportes_del_catalogo_generan_pdf(self):
        from .catalogo import REPORTES

        for clave, definicion in REPORTES.items():
            with self.subTest(reporte=clave):
                _, contenido, _, _ = construir_reporte(
                    self.admin, definicion, definicion.columnas, "PDF", self.filtros
                )
                self.assertTrue(contenido.startswith(b"%PDF"))

    def test_todos_los_reportes_generan_excel_con_columnas_ampliadas(self):
        """Con datos reales en cada modulo (entregables+evaluacion, gastos
        comprometidos/ejecutados, empresa con contactos y documentos), las
        columnas nuevas de cada reporte no deben tronar ni hacer N+1 raro."""
        from .catalogo import REPORTES

        self._sembrar_datos_ricos()
        for clave, definicion in REPORTES.items():
            with self.subTest(reporte=clave):
                _, contenido, nombre, tipo = construir_reporte(
                    self.admin, definicion, definicion.columnas, "EXCEL", self.filtros
                )
                self.assertTrue(contenido.startswith(b"PK"))
                self.assertTrue(nombre.endswith(".xlsx"))


class AccesoAReportesTests(TestCase):
    def setUp(self):
        self.admin = Usuario.objects.create_superuser(
            username="admin-acc", password="x", email="a@ean.edu.co"
        )
        self.estudiante = Usuario.objects.create_user(
            username="est-acc", password="x", rol="ESTUDIANTE"
        )

    def test_admin_abre_el_constructor(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("reportes:constructor"))
        self.assertEqual(response.status_code, 200)

    def test_estudiante_no_accede_a_reportes(self):
        self.client.force_login(self.estudiante)
        response = self.client.get(reverse("reportes:constructor"))
        self.assertEqual(response.status_code, 302)

    def test_no_se_descarga_el_reporte_de_otro(self):
        otro = Usuario.objects.create_user(username="profe-acc", password="x", rol="PROFESOR")
        reporte = ReporteGenerado.objects.create(usuario=self.admin, formato="PDF", filtros={})
        self.client.force_login(otro)
        response = self.client.get(reverse("reportes:descargar", args=[reporte.pk]))
        self.assertRedirects(response, reverse("reportes:historial"))

    def test_generar_desde_la_vista_devuelve_el_archivo(self):
        self.client.force_login(self.admin)
        response = self.client.post(reverse("reportes:constructor"), {
            "reporte": "retos",
            "formato": "EXCEL",
            "periodo": "todo",
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(ReporteGenerado.objects.count(), 1)


class DatosNuevosEnLosReportesTests(TestCase):
    """Los campos que trajo la integracion con la rama de Caicedo se exportan con su valor real."""

    def setUp(self):
        from apps.empresas.models import Empresa

        self.admin = Usuario.objects.create_superuser(username="admin-nuevos", password="x", email="a@ean.edu.co")
        self.usuario = Usuario.objects.create_user(username="empresa-nuevos", password="x", rol="EMPRESA")
        self.empresa = Empresa.objects.create(
            usuario=self.usuario, nit="900111000", razon_social="Epsilon SAS",
            tipo_convocatoria="AMBOS", area_responsable="Innovacion", sitio_web="https://epsilon.example",
            participacion_bloqueada=True, motivo_bloqueo="Pendiente de revision juridica",
        )
        self.reto = Reto.objects.create(
            empresa=self.usuario, titulo="Reto con registro completo", estado="aprobado",
            nombre_corto="RCRC", uso_esperado_resultados="Prototipo en produccion",
            insumos_inicio="Datos historicos", insumos_cierre="Informe final",
        )
        self.filtros = {"desde": None, "hasta": None}

    def _filas(self, clave):
        from io import BytesIO

        from openpyxl import load_workbook

        definicion = obtener(clave)
        _, contenido, _, _ = construir_reporte(self.admin, definicion, definicion.columnas, "EXCEL", self.filtros)
        filas = list(load_workbook(BytesIO(contenido)).active.iter_rows(values_only=True))
        encabezado = next(f for f in filas if f and f[0] and f[1] and f[0] != definicion.nombre and "Generado" not in str(f[0]))
        return encabezado, filas

    def _valor(self, clave, marca, columna):
        encabezado, filas = self._filas(clave)
        fila = next(f for f in filas if marca in [str(c) for c in f if c])
        return fila[encabezado.index(columna)]

    def test_empresas_exporta_registro_bloqueo_y_pendientes_documentales(self):
        marca = "Epsilon SAS"
        self.assertEqual(self._valor("empresas", marca, "Tipo de convocatoria"), "Reto y hackathon")
        self.assertEqual(self._valor("empresas", marca, "Area responsable"), "Innovacion")
        self.assertEqual(self._valor("empresas", marca, "Sitio web"), "https://epsilon.example")
        self.assertEqual(self._valor("empresas", marca, "Participacion bloqueada"), "Si")
        self.assertEqual(self._valor("empresas", marca, "Motivo del bloqueo"), "Pendiente de revision juridica")
        self.assertEqual(self._valor("empresas", marca, "Puede participar"), "No")
        self.assertIn("Falta cargarlo", self._valor("empresas", marca, "Pendientes para participar"))

    def test_retos_exporta_registro_e_impacto_del_cierre(self):
        from apps.cierre.models import CierreReto

        CierreReto.objects.create(reto=self.reto, impacto="SI")
        marca = "Reto con registro completo"
        self.assertEqual(self._valor("retos", marca, "Nombre corto"), "RCRC")
        self.assertEqual(self._valor("retos", marca, "Uso esperado de los resultados"), "Prototipo en produccion")
        self.assertEqual(self._valor("retos", marca, "Insumos al inicio"), "Datos historicos")
        self.assertEqual(self._valor("retos", marca, "Insumos al cierre"), "Informe final")
        self.assertEqual(self._valor("retos", marca, "La empresa aplico la solucion"), "Sí, la aplicó")

    def test_participacion_y_evaluacion_exportan_los_campos_nuevos(self):
        from apps.evaluacion.models import Entregable

        estudiante = Usuario.objects.create_user(username="est-nuevos", password="x", rol="ESTUDIANTE", first_name="Eva", last_name="Nueva")
        Postulacion.objects.create(reto=self.reto, estudiante=estudiante, estado="PENDIENTE", acepta_terminos=True)
        Entregable.objects.create(
            reto=self.reto, estudiante=estudiante, titulo="Avance publicado", estado="CALIFICADO",
            nota=4, calificacion_publicada=True,
        )
        self.assertEqual(self._valor("participacion", "Eva Nueva", "Acepto los terminos"), "Si")
        self.assertEqual(self._valor("participacion", "Eva Nueva", "Adjunto hoja de vida"), "No")
        self.assertEqual(self._valor("evaluacion", "Avance publicado", "Calificacion publicada al estudiante"), "Si")


class ReportesConDatosHostilesTests(TestCase):
    """Los reportes llevan texto libre de la base: no debe romper ni el PDF ni el Excel."""

    def setUp(self):
        from apps.presupuesto.models import Gasto, Presupuesto

        self.admin = Usuario.objects.create_superuser(username="admin-hostil-rep", password="x", email="a@ean.edu.co")
        empresa = Usuario.objects.create_user(username="empresa-hostil-rep", password="x", rol="EMPRESA")
        self.reto = Reto.objects.create(empresa=empresa, titulo="Reto <b & =cmd", estado="aprobado",
                                        descripcion="con\x01control y <i")
        presupuesto = Presupuesto.objects.create(reto=self.reto, monto_total=100)
        # Una categoria que ya no existe en el catalogo (dato antiguo o cargado desde el admin).
        Gasto.objects.create(presupuesto=presupuesto, categoria="VIEJA", monto=10)
        self.filtros = {"desde": None, "hasta": None}

    def test_todos_los_reportes_se_generan_en_pdf_y_excel(self):
        from .catalogo import REPORTES

        for clave, definicion in REPORTES.items():
            for formato, firma in (("PDF", b"%PDF"), ("EXCEL", b"PK")):
                with self.subTest(reporte=clave, formato=formato):
                    _, contenido, _, _ = construir_reporte(
                        self.admin, definicion, definicion.columnas, formato, self.filtros)
                    self.assertTrue(contenido.startswith(firma))


class VistaPreviaYEnvioHU06Tests(TestCase):
    """HU06: vista previa en pantalla y envio del reporte por correo."""

    def setUp(self):
        self.admin = Usuario.objects.create_superuser(username="admin-hu06", password="x", email="admin@ean.edu.co")
        empresa = Usuario.objects.create_user(username="emp-hu06", password="x", rol="EMPRESA")
        for i in range(3):
            Reto.objects.create(empresa=empresa, titulo=f"Reto previo {i}", estado="aprobado")
        self.client.force_login(self.admin)

    def _post(self, **extra):
        return self.client.post(reverse("reportes:constructor"), {
            "reporte": "retos", "formato": "PDF", "periodo": "todo", **extra,
        })

    def test_vista_previa_muestra_filas_sin_generar_archivo(self):
        respuesta = self._post(accion="previsualizar")
        self.assertEqual(respuesta.status_code, 200)
        previa = respuesta.context["vista_previa"]
        self.assertEqual(previa["total"], 3)
        self.assertEqual(len(previa["filas"]), 3)
        self.assertContains(respuesta, "Reto previo 0")
        self.assertFalse(ReporteGenerado.objects.exists())

    def test_enviar_por_correo_adjunta_el_archivo_y_lo_registra(self):
        from django.core import mail

        respuesta = self._post(accion="enviar", destinatarios="jefe@ean.edu.co; otra@ean.edu.co")
        self.assertRedirects(respuesta, reverse("reportes:historial"))
        self.assertEqual(len(mail.outbox), 1)
        correo = mail.outbox[0]
        self.assertEqual(correo.to, ["jefe@ean.edu.co", "otra@ean.edu.co"])
        self.assertEqual(correo.attachments[0][2], "application/pdf")
        registro = ReporteGenerado.objects.get()
        self.assertTrue(registro.enviado_por_email)
        self.assertEqual(registro.filtros["enviado_a"], ["jefe@ean.edu.co", "otra@ean.edu.co"])

    def test_sin_destinatarios_se_envia_al_propio_usuario(self):
        from django.core import mail

        self._post(accion="enviar")
        self.assertEqual(mail.outbox[0].to, ["admin@ean.edu.co"])

    def test_correo_invalido_no_genera_nada(self):
        from django.core import mail

        respuesta = self._post(accion="enviar", destinatarios="no-es-correo")
        self.assertEqual(respuesta.status_code, 200)
        self.assertFalse(ReporteGenerado.objects.exists())
        self.assertEqual(len(mail.outbox), 0)

    def test_fallo_de_envio_conserva_el_reporte_y_lo_marca(self):
        from unittest import mock

        with mock.patch("django.core.mail.EmailMessage.send", side_effect=OSError("SMTP caido")):
            self._post(accion="enviar")
        registro = ReporteGenerado.objects.get()
        self.assertFalse(registro.enviado_por_email)
        self.assertIn("SMTP caido", registro.filtros["error_envio"])
        historial = self.client.get(reverse("reportes:historial"))
        self.assertContains(historial, "Fallido")
