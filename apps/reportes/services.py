"""Generacion y persistencia de reportes (HU06).

PDF con reportlab y Excel con openpyxl: ambas son puro Python y traen ruedas
para Windows, a diferencia de WeasyPrint, que exige GTK/Pango nativos.
"""

import io
from datetime import date
from xml.sax.saxutils import escape

from django.core.files.base import ContentFile
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .models import ReporteGenerado

# Misma paleta de marca que templates/base.html (variables --ean-*), para que
# el PDF exportado se vea como una extension de la plataforma, no un genero.
VERDE_EAN = colors.HexColor("#3bac52")
AZUL_EAN = colors.HexColor("#005e85")
NEGRO_EAN = colors.HexColor("#0a0a0a")
GRIS_TEXTO = colors.HexColor("#6b7280")
GRIS_SUAVE = colors.HexColor("#f3f4f6")

MAX_FILAS_PDF = 2000


def _filas(definicion, queryset, columnas):
    for objeto in queryset:
        yield [columna.obtener(objeto) for columna in columnas]


def _describir_filtros(filtros):
    """Texto legible con los filtros aplicados, para el pie del reporte."""
    partes = []
    etiquetas = {
        "desde": "Desde", "hasta": "Hasta", "estado": "Estado",
        "estado_postulacion": "Estado de postulacion", "estado_empresa": "Estado de empresa",
        "facultad": "Facultad", "programa": "Programa", "empresa": "Empresa",
    }
    for clave, etiqueta in etiquetas.items():
        valor = filtros.get(clave)
        if valor:
            if isinstance(valor, date):
                valor = valor.strftime("%d/%m/%Y")
            partes.append(f"{etiqueta}: {valor}")
    return " | ".join(partes) or "Sin filtros adicionales"


# ------------------------------------------------------------------ EXCEL
def agregar_fila(hoja, valores):
    """Escribe una fila sin que un texto de la base rompa el archivo ni se vuelva formula.

    openpyxl convierte en formula todo texto que empiece por "=" y lanza error con
    caracteres de control; los reportes llevan texto libre (descripciones, titulos).
    """
    hoja.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in valores])
    for celda in hoja[hoja.max_row]:
        if isinstance(celda.value, str) and celda.value.startswith("="):
            celda.data_type = "s"


def generar_excel(definicion, queryset, columnas, filtros):
    libro = Workbook()
    hoja = libro.active
    hoja.title = definicion.nombre[:31]

    # Misma paleta de marca que el PDF y la plataforma (verde EAN), en vez
    # del gris generico que tenia antes.
    encabezado_fondo = PatternFill("solid", fgColor="3BAC52")
    encabezado_fuente = Font(color="FFFFFF", bold=True, size=11)

    hoja.append(["retosean · Universidad EAN"])
    hoja["A1"].font = Font(bold=True, size=11, color="3BAC52")
    hoja.append([definicion.nombre])
    hoja["A2"].font = Font(bold=True, size=14, color="0A0A0A")
    hoja.append([f"Generado el {timezone.localtime():%d/%m/%Y %H:%M}"])
    agregar_fila(hoja, [_describir_filtros(filtros)])
    hoja.append([])

    fila_encabezado = hoja.max_row + 1
    hoja.append([columna.etiqueta for columna in columnas])
    for indice in range(1, len(columnas) + 1):
        celda = hoja.cell(row=fila_encabezado, column=indice)
        celda.fill = encabezado_fondo
        celda.font = encabezado_fuente
        celda.alignment = Alignment(vertical="center", wrap_text=True)

    total = 0
    for fila in _filas(definicion, queryset, columnas):
        agregar_fila(hoja, fila)
        total += 1

    for indice, columna in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = columna.ancho
    hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)

    hoja.append([])
    hoja.append([f"Total de registros: {total}"])

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue(), total


# -------------------------------------------------------------------- PDF
def encabezado_de_marca(canvas, documento):
    """Misma franja de marca en cada pagina: acento verde arriba, wordmark +
    Universidad EAN, y pie con paginacion. Reemplaza la marca de agua generica
    (gris/azul, sin identidad) que tenia el PDF antes."""
    canvas.saveState()
    ancho, alto = documento.pagesize

    # Franja de acento superior, como el borde del topbar de la plataforma.
    canvas.setFillColor(VERDE_EAN)
    canvas.rect(0, alto - 3 * mm, ancho, 3 * mm, stroke=0, fill=1)

    canvas.setFont("Helvetica-Bold", 10)
    canvas.setFillColor(NEGRO_EAN)
    canvas.drawString(15 * mm, alto - 10 * mm, "retos")
    ancho_retos = canvas.stringWidth("retos", "Helvetica-Bold", 10)
    canvas.setFillColor(VERDE_EAN)
    canvas.drawString(15 * mm + ancho_retos, alto - 10 * mm, "ean")

    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(GRIS_TEXTO)
    canvas.drawRightString(ancho - 15 * mm, alto - 10 * mm, "Universidad EAN")

    canvas.setStrokeColor(AZUL_EAN)
    canvas.setLineWidth(0.6)
    canvas.line(15 * mm, alto - 13 * mm, ancho - 15 * mm, alto - 13 * mm)

    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(GRIS_TEXTO)
    canvas.drawString(15 * mm, 10 * mm, "RetosEAN - Universidad EAN")
    canvas.drawRightString(ancho - 15 * mm, 10 * mm, f"Pagina {canvas.getPageNumber()}")
    canvas.restoreState()


def generar_pdf(definicion, queryset, columnas, filtros):
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=22 * mm, bottomMargin=18 * mm,
        title=definicion.nombre,
        author="RetosEAN - Universidad EAN",
    )

    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        "TituloReporte", parent=estilos["Title"],
        fontSize=16, textColor=NEGRO_EAN, spaceAfter=4, alignment=0,
    )
    estilo_meta = ParagraphStyle(
        "MetaReporte", parent=estilos["Normal"],
        fontSize=8.5, textColor=GRIS_TEXTO, spaceAfter=2,
    )
    estilo_celda = ParagraphStyle(
        "Celda", parent=estilos["Normal"], fontSize=7.5, leading=9.5,
    )
    estilo_encabezado = ParagraphStyle(
        "CeldaEncabezado", parent=estilos["Normal"],
        fontSize=8, leading=10, textColor=colors.white, fontName="Helvetica-Bold",
    )

    elementos = [
        Paragraph(definicion.nombre, estilo_titulo),
        Paragraph(definicion.descripcion, estilo_meta),
        Paragraph(f"Generado el {timezone.localtime():%d/%m/%Y %H:%M}", estilo_meta),
        Paragraph(escape(_describir_filtros(filtros)), estilo_meta),
        Spacer(1, 8),
    ]

    datos = [[Paragraph(c.etiqueta, estilo_encabezado) for c in columnas]]
    total = 0
    for fila in _filas(definicion, queryset, columnas):
        if total >= MAX_FILAS_PDF:
            break
        datos.append([Paragraph(escape(str(valor)), estilo_celda) for valor in fila])
        total += 1

    if total == 0:
        elementos.append(Paragraph("No hay registros que cumplan los filtros seleccionados.", estilos["Normal"]))
    else:
        ancho_util = documento.width
        pesos = [c.ancho for c in columnas]
        suma = sum(pesos) or 1
        anchos = [ancho_util * peso / suma for peso in pesos]

        tabla = Table(datos, colWidths=anchos, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), VERDE_EAN),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d5db")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, GRIS_SUAVE]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elementos.append(tabla)
        elementos.append(Spacer(1, 8))
        elementos.append(Paragraph(
            f"Total de registros: {total}"
            + (f" (truncado a {MAX_FILAS_PDF}; exporta a Excel para el listado completo)"
               if total >= MAX_FILAS_PDF else ""),
            ParagraphStyle("Total", parent=estilo_meta, alignment=TA_RIGHT),
        ))

    documento.build(elementos, onFirstPage=encabezado_de_marca, onLaterPages=encabezado_de_marca)
    return buffer.getvalue(), total


# ------------------------------------------------------------ ORQUESTACION
def _filtros_serializables(filtros):
    resultado = {}
    for clave, valor in filtros.items():
        if valor in (None, "", []):
            continue
        resultado[clave] = valor.isoformat() if isinstance(valor, date) else str(valor)
    return resultado


FILAS_VISTA_PREVIA = 50


def vista_previa(definicion, columnas, filtros):
    """Primeras filas y total del reporte, sin generar archivo ni registro (HU06)."""
    queryset = definicion.construir_queryset(filtros)
    total = queryset.count()
    filas = list(_filas(definicion, queryset[:FILAS_VISTA_PREVIA], columnas))
    return {
        "encabezados": [c.etiqueta for c in columnas],
        "filas": filas,
        "total": total,
        "mostradas": len(filas),
        "filtros": _describir_filtros(filtros),
    }


def enviar_por_correo(registro, contenido, nombre, tipo_mime, destinatarios, definicion):
    """Envia el reporte como adjunto y deja constancia en el historial (HU06).

    Devuelve None si se envio, o el mensaje de error; el reporte generado se
    conserva en el historial aunque el envio falle.
    """
    from django.conf import settings
    from django.core.mail import EmailMessage

    autor = registro.usuario.get_full_name() or registro.usuario.username
    correo = EmailMessage(
        subject=f"[RetosEAN] {definicion.nombre}",
        body=(
            f"{autor} te envia el reporte \"{definicion.nombre}\" generado en RetosEAN "
            f"el {timezone.localtime(registro.creado_en):%d/%m/%Y %H:%M}.\n\n"
            f"Registros: {registro.filtros.get('total_registros', 0)}.\n"
            "El archivo va adjunto."
        ),
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=destinatarios,
    )
    correo.attach(nombre, contenido, tipo_mime)
    try:
        correo.send(fail_silently=False)
    except Exception as error:  # SMTP, DNS o credenciales: se informa sin perder el reporte.
        registro.filtros = {**registro.filtros, "error_envio": str(error)[:300]}
        registro.save(update_fields=["filtros"])
        return str(error)
    registro.enviado_por_email = True
    registro.filtros = {**registro.filtros, "enviado_a": destinatarios}
    registro.save(update_fields=["enviado_por_email", "filtros"])
    return None


def construir_reporte(usuario, definicion, columnas, formato, filtros):
    """Genera el archivo y lo deja registrado en ReporteGenerado."""
    queryset = definicion.construir_queryset(filtros)

    if formato == "EXCEL":
        contenido, total = generar_excel(definicion, queryset, columnas, filtros)
        extension, tipo_mime = "xlsx", (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    else:
        contenido, total = generar_pdf(definicion, queryset, columnas, filtros)
        extension, tipo_mime = "pdf", "application/pdf"

    marca = timezone.localtime().strftime("%Y%m%d-%H%M%S")
    nombre = f"{definicion.clave}-{marca}.{extension}"

    registro = ReporteGenerado(
        usuario=usuario,
        formato=formato,
        filtros={
            "reporte": definicion.clave,
            "columnas": [c.clave for c in columnas],
            "total_registros": total,
            **_filtros_serializables(filtros),
        },
    )
    registro.archivo.save(nombre, ContentFile(contenido), save=True)
    return registro, contenido, nombre, tipo_mime
