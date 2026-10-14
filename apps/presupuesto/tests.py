from decimal import Decimal

from django.test import TestCase
from django.urls import reverse

from apps.notificaciones.models import Notificacion
from apps.retos.models import Reto
from apps.usuarios.models import Usuario

from .models import Gasto, Presupuesto
from .services import ejecutar_gasto, registrar_gasto, revisar_umbral


class PresupuestoModeloTests(TestCase):
    def setUp(self):
        empresa = Usuario.objects.create_user(username="empresa-pre", password="x", rol="EMPRESA")
        self.reto = Reto.objects.create(empresa=empresa, titulo="Reto con plata", estado="aprobado")
        self.presupuesto = Presupuesto.objects.create(reto=self.reto, monto_total=Decimal("1000.00"))

    def test_sin_gastos_el_ejecutado_es_decimal_cero(self):
        ejecutado = self.presupuesto.monto_ejecutado
        self.assertEqual(ejecutado, Decimal("0.00"))
        self.assertIsInstance(ejecutado, Decimal)

    def test_calcula_ejecutado_disponible_y_porcentaje(self):
        Gasto.objects.create(presupuesto=self.presupuesto, categoria="MATERIALES", monto=Decimal("250.00"))
        Gasto.objects.create(presupuesto=self.presupuesto, categoria="REFRIGERIOS", monto=Decimal("150.00"))
        self.assertEqual(self.presupuesto.monto_ejecutado, Decimal("400.00"))
        self.assertEqual(self.presupuesto.monto_disponible, Decimal("600.00"))
        self.assertEqual(self.presupuesto.porcentaje_ejecutado, Decimal("40"))

    def test_presupuesto_en_cero_no_divide_por_cero(self):
        vacio = Presupuesto.objects.create(
            reto=Reto.objects.create(
                empresa=self.reto.empresa, titulo="Sin plata", estado="aprobado"
            ),
            monto_total=Decimal("0.00"),
        )
        self.assertEqual(vacio.porcentaje_ejecutado, Decimal("0.00"))


class ComprometidoVsEjecutadoTests(TestCase):
    def setUp(self):
        empresa = Usuario.objects.create_user(username="empresa-ce", password="x", rol="EMPRESA")
        self.reto = Reto.objects.create(empresa=empresa, titulo="Reto con reservas", estado="aprobado")
        self.presupuesto = Presupuesto.objects.create(reto=self.reto, monto_total=Decimal("1000.00"))

    def test_gasto_sin_estado_explicito_cuenta_como_ejecutado(self):
        # Compatibilidad con gastos historicos, previos a este campo.
        Gasto.objects.create(presupuesto=self.presupuesto, categoria="OTROS", monto=Decimal("100.00"))
        self.assertEqual(self.presupuesto.monto_ejecutado, Decimal("100.00"))
        self.assertEqual(self.presupuesto.monto_comprometido, Decimal("0.00"))

    def test_comprometido_y_ejecutado_se_suman_por_separado(self):
        Gasto.objects.create(
            presupuesto=self.presupuesto, categoria="MATERIALES", monto=Decimal("300.00"),
            estado=Gasto.COMPROMETIDO,
        )
        Gasto.objects.create(
            presupuesto=self.presupuesto, categoria="REFRIGERIOS", monto=Decimal("200.00"),
            estado=Gasto.EJECUTADO,
        )
        self.assertEqual(self.presupuesto.monto_comprometido, Decimal("300.00"))
        self.assertEqual(self.presupuesto.monto_ejecutado, Decimal("200.00"))
        # Disponible y porcentaje se calculan sobre comprometido + ejecutado.
        self.assertEqual(self.presupuesto.monto_disponible, Decimal("500.00"))
        self.assertEqual(self.presupuesto.porcentaje_ejecutado, Decimal("50"))

    def test_ejecutar_gasto_lo_mueve_de_comprometido_a_ejecutado(self):
        gasto = Gasto.objects.create(
            presupuesto=self.presupuesto, categoria="OTROS", monto=Decimal("400.00"),
            estado=Gasto.COMPROMETIDO,
        )
        ejecutar_gasto(gasto)
        gasto.refresh_from_db()
        self.assertEqual(gasto.estado, Gasto.EJECUTADO)
        self.assertIsNotNone(gasto.ejecutado_en)
        self.assertEqual(self.presupuesto.monto_comprometido, Decimal("0.00"))
        self.assertEqual(self.presupuesto.monto_ejecutado, Decimal("400.00"))

    def test_ejecutar_gasto_ya_ejecutado_no_hace_nada(self):
        gasto = Gasto.objects.create(
            presupuesto=self.presupuesto, categoria="OTROS", monto=Decimal("100.00"),
            estado=Gasto.EJECUTADO,
        )
        resultado = ejecutar_gasto(gasto)
        self.assertEqual(resultado.estado, Gasto.EJECUTADO)

    def test_comprometido_por_si_solo_puede_disparar_la_alerta_del_80(self):
        # Reservar (no gastar todavia) el 85% del presupuesto ya debe alertar.
        registrar_gasto(
            self.presupuesto,
            Gasto(categoria="OTROS", monto=Decimal("850.00"), estado=Gasto.COMPROMETIDO),
            self.reto.empresa,
        )
        self.presupuesto.refresh_from_db()
        self.assertTrue(self.presupuesto.alerta_80_enviada)


class AlertaDelOchentaTests(TestCase):
    def setUp(self):
        self.admin = Usuario.objects.create_superuser(
            username="admin-pre", password="x", email="admin@ean.edu.co"
        )
        self.empresa = Usuario.objects.create_user(
            username="empresa-alerta", password="x", rol="EMPRESA"
        )
        self.reto = Reto.objects.create(
            empresa=self.empresa, titulo="Reto vigilado", estado="en_curso"
        )
        self.presupuesto = Presupuesto.objects.create(reto=self.reto, monto_total=Decimal("1000.00"))

    def _gasto(self, monto):
        return registrar_gasto(
            self.presupuesto,
            Gasto(categoria="OTROS", monto=Decimal(monto)),
            self.admin,
        )

    def test_por_debajo_del_umbral_no_alerta(self):
        self._gasto("500.00")
        self.presupuesto.refresh_from_db()
        self.assertFalse(self.presupuesto.alerta_80_enviada)
        self.assertEqual(Notificacion.objects.filter(evento="PRESUPUESTO_80").count(), 0)

    def test_al_superar_el_80_notifica_a_admin_y_empresa(self):
        self._gasto("850.00")
        self.presupuesto.refresh_from_db()
        self.assertTrue(self.presupuesto.alerta_80_enviada)
        destinatarios = set(
            Notificacion.objects.filter(evento="PRESUPUESTO_80").values_list("usuario_id", flat=True)
        )
        self.assertIn(self.admin.pk, destinatarios)
        self.assertIn(self.empresa.pk, destinatarios)

    def test_la_alerta_no_se_repite_en_cada_gasto(self):
        self._gasto("850.00")
        self._gasto("50.00")
        self._gasto("30.00")
        self.assertEqual(Notificacion.objects.filter(evento="PRESUPUESTO_80").count(), 2)

    def test_ampliar_el_presupuesto_rearma_la_alerta(self):
        self._gasto("850.00")
        self.presupuesto.refresh_from_db()
        self.assertTrue(self.presupuesto.alerta_80_enviada)

        self.presupuesto.monto_total = Decimal("5000.00")
        self.presupuesto.save(update_fields=["monto_total"])
        revisar_umbral(self.presupuesto)
        self.presupuesto.refresh_from_db()
        self.assertFalse(self.presupuesto.alerta_80_enviada)


class AccesoAPresupuestoTests(TestCase):
    def setUp(self):
        self.admin = Usuario.objects.create_superuser(
            username="admin-acceso", password="x", email="a@ean.edu.co"
        )
        self.empresa = Usuario.objects.create_user(
            username="empresa-acceso", password="x", rol="EMPRESA"
        )
        self.otra_empresa = Usuario.objects.create_user(
            username="empresa-ajena", password="x", rol="EMPRESA"
        )
        self.reto = Reto.objects.create(
            empresa=self.empresa, titulo="Reto propio", estado="aprobado"
        )
        Presupuesto.objects.create(reto=self.reto, monto_total=Decimal("500.00"))

    def test_admin_ve_el_panel(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("presupuesto:panel")).status_code, 200)

    def test_empresa_ve_el_presupuesto_de_su_reto(self):
        self.client.force_login(self.empresa)
        response = self.client.get(reverse("presupuesto:detalle", args=[self.reto.pk]))
        self.assertEqual(response.status_code, 200)

    def test_empresa_no_ve_el_presupuesto_ajeno(self):
        self.client.force_login(self.otra_empresa)
        response = self.client.get(reverse("presupuesto:detalle", args=[self.reto.pk]))
        self.assertEqual(response.status_code, 302)

    def test_empresa_no_puede_registrar_gastos(self):
        self.client.force_login(self.empresa)
        response = self.client.post(
            reverse("presupuesto:agregar_gasto", args=[self.reto.pk]),
            {"categoria": "OTROS", "monto": "100"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Gasto.objects.count(), 0)


class ExportacionPresupuestoTests(TestCase):
    """HU07: reporte de gastos en Excel y PDF."""

    def setUp(self):
        self.empresa = Usuario.objects.create_user(username="empresa-exp", password="x", rol="EMPRESA")
        self.admin = Usuario.objects.create_user(username="admin-exp", password="x", rol="ADMIN")
        self.reto = Reto.objects.create(empresa=self.empresa, titulo="Reto exportable", estado="en_curso")
        self.presupuesto = Presupuesto.objects.create(reto=self.reto, monto_total=Decimal("1000.00"))
        Gasto.objects.create(presupuesto=self.presupuesto, categoria="MATERIALES", monto=Decimal("250.00"),
                             descripcion="Cartulinas")

    def test_excel_trae_resumen_y_detalle(self):
        import io

        from openpyxl import load_workbook

        self.client.force_login(self.admin)
        response = self.client.get(reverse("presupuesto:exportar", args=[self.reto.pk, "xlsx"]))
        self.assertEqual(response.status_code, 200)
        valores = [c for fila in load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True) for c in fila]
        self.assertIn("Cartulinas", valores)
        self.assertIn(250.0, valores)
        self.assertIn(1000.0, valores)

    def test_pdf_se_genera(self):
        self.client.force_login(self.empresa)
        response = self.client.get(reverse("presupuesto:exportar", args=[self.reto.pk, "pdf"]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_otra_empresa_no_exporta(self):
        otra = Usuario.objects.create_user(username="empresa-ajena", password="x", rol="EMPRESA")
        self.client.force_login(otra)
        response = self.client.get(reverse("presupuesto:exportar", args=[self.reto.pk, "pdf"]))
        self.assertNotEqual(response.status_code, 200)

    def test_formato_desconocido_es_404(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("presupuesto:exportar", args=[self.reto.pk, "doc"])).status_code, 404)


class ExportacionConTextoHostilTests(TestCase):
    """Un texto libre con simbolos (descripciones, titulos) no debe tumbar ni alterar la descarga."""

    def setUp(self):
        self.admin = Usuario.objects.create_user(username="admin-hostil", password="x", rol="ADMIN")
        empresa = Usuario.objects.create_user(username="empresa-hostil", password="x", rol="EMPRESA")
        self.reto = Reto.objects.create(empresa=empresa, titulo="Reto <b> & cia", estado="en_curso")
        self.presupuesto = Presupuesto.objects.create(reto=self.reto, monto_total=Decimal("1000.00"))
        for descripcion in ("Tinta <b", "Pago <proveedor> & cia", "=1+1", "con\x01control"):
            Gasto.objects.create(presupuesto=self.presupuesto, categoria="OTROS",
                                 monto=Decimal("10.00"), descripcion=descripcion)
        self.client.force_login(self.admin)

    def _descargar(self, formato):
        return self.client.get(reverse("presupuesto:exportar", args=[self.reto.pk, formato]))

    def test_el_pdf_se_genera_con_simbolos_en_el_texto(self):
        respuesta = self._descargar("pdf")
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(respuesta.content.startswith(b"%PDF"))

    def test_el_excel_no_convierte_texto_en_formula_ni_falla_con_caracteres_de_control(self):
        import io

        from openpyxl import load_workbook

        respuesta = self._descargar("xlsx")
        self.assertEqual(respuesta.status_code, 200)
        celdas = [c for fila in load_workbook(io.BytesIO(respuesta.content)).active.iter_rows() for c in fila]
        formula = next(c for c in celdas if c.value == "=1+1")
        self.assertNotEqual(formula.data_type, "f")  # es texto, no una formula que Excel evalue
        self.assertIn("concontrol", [c.value for c in celdas])  # el caracter de control se limpia
