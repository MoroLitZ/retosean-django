"""Registro de gastos y alerta de ejecucion presupuestal (HU07)."""

from decimal import Decimal

from django.db import transaction
from django.urls import reverse
from django.utils import timezone

from apps.notificaciones.services import notificar, notificar_admins

from .models import Gasto

UMBRAL_ALERTA = Decimal("80")


@transaction.atomic
def registrar_gasto(presupuesto, gasto, usuario):
    """Guarda el gasto (comprometido o ejecutado) y revisa el umbral de alerta.

    El umbral se calcula sobre comprometido + ejecutado, asi que reservar
    presupuesto ya puede disparar la alerta, no solo gastarlo de verdad.
    """
    gasto.presupuesto = presupuesto
    gasto.registrado_por = usuario
    if gasto.estado == Gasto.EJECUTADO and gasto.ejecutado_en is None:
        gasto.ejecutado_en = timezone.now()
    gasto.save()
    revisar_umbral(presupuesto)
    return gasto


@transaction.atomic
def ejecutar_gasto(gasto):
    """Pasa un gasto de COMPROMETIDO a EJECUTADO (ya se pago de verdad)."""
    if gasto.estado != Gasto.COMPROMETIDO:
        return gasto
    gasto.estado = Gasto.EJECUTADO
    gasto.ejecutado_en = timezone.now()
    gasto.save(update_fields=["estado", "ejecutado_en"])
    revisar_umbral(gasto.presupuesto)
    return gasto


def revisar_umbral(presupuesto):
    """Notifica una sola vez cuando comprometido + ejecutado supera el umbral.

    Se calcula sobre lo asignado (comprometido + ejecutado), no solo lo ya
    pagado: reservar presupuesto ya reduce lo disponible de verdad, asi que
    debe poder disparar la alerta igual que un gasto pagado.
    """
    porcentaje = presupuesto.porcentaje_ejecutado

    if porcentaje < UMBRAL_ALERTA:
        # Si el presupuesto se amplia o se borra/reduce un gasto, se rearma la alerta.
        if presupuesto.alerta_80_enviada:
            presupuesto.alerta_80_enviada = False
            presupuesto.save(update_fields=["alerta_80_enviada"])
        return False

    if presupuesto.alerta_80_enviada:
        return False

    reto = presupuesto.reto
    mensaje = (
        f"El presupuesto del reto '{reto.titulo}' va en {porcentaje:.1f}% asignado "
        f"(comprometido: {presupuesto.monto_comprometido}, ejecutado: {presupuesto.monto_ejecutado}, "
        f"de {presupuesto.monto_total}). Disponible: {presupuesto.monto_disponible}."
    )
    link = reverse("presupuesto:detalle", kwargs={"reto_id": reto.pk})

    notificar_admins(
        "PRESUPUESTO_80", mensaje=mensaje, link=link,
        clave_dedupe=f"presupuesto:{presupuesto.pk}:80",
    )
    notificar(
        reto.empresa, "PRESUPUESTO_80", mensaje=mensaje, link=link,
        clave_dedupe=f"presupuesto:{presupuesto.pk}:80",
    )

    presupuesto.alerta_80_enviada = True
    presupuesto.save(update_fields=["alerta_80_enviada"])
    return True


# ------------------------------------------------------------- EXPORTACION
def _resumen_por_categoria(presupuesto):
    """[(categoria, comprometido, ejecutado)] con todas las categorias, aunque no tengan gastos."""
    from django.db.models import Sum

    from .models import Gasto

    cero = Decimal("0")
    totales = {
        (fila["categoria"], fila["estado"]): fila["total"]
        for fila in presupuesto.gastos.values("categoria", "estado").annotate(total=Sum("monto"))
    }
    return [
        (etiqueta, totales.get((valor, Gasto.COMPROMETIDO), cero), totales.get((valor, Gasto.EJECUTADO), cero))
        for valor, etiqueta in Gasto.CATEGORIAS
    ]


def _gastos_ordenados(presupuesto):
    return presupuesto.gastos.select_related("registrado_por").order_by("creado_en")


def _resumen_general(presupuesto):
    """[(etiqueta, valor)] del encabezado de ambos reportes; el porcentaje es el unico no monetario."""
    return [
        ("Asignado", presupuesto.monto_total),
        ("Comprometido", presupuesto.monto_comprometido),
        ("Ejecutado", presupuesto.monto_ejecutado),
        ("Disponible", presupuesto.monto_disponible),
        ("% asignado", round(presupuesto.porcentaje_ejecutado, 2)),
    ]


def exportar_excel(presupuesto):
    """Reporte de gastos del reto en Excel: resumen, totales por categoria y detalle (HU07)."""
    import io

    from django.utils import timezone
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    from apps.reportes.services import agregar_fila

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Gastos"
    relleno = PatternFill("solid", fgColor="3BAC52")  # verde EAN, igual que los reportes
    blanca = Font(color="FFFFFF", bold=True)

    def encabezado(*valores):
        hoja.append(list(valores))
        for celda in hoja[hoja.max_row]:
            celda.fill = relleno
            celda.font = blanca

    hoja.append(["retosean · Universidad EAN"])
    hoja["A1"].font = Font(bold=True, size=11, color="3BAC52")
    agregar_fila(hoja, [f"Presupuesto del reto: {presupuesto.reto.titulo}"])
    hoja["A2"].font = Font(bold=True, size=14)
    hoja.append([f"Generado el {timezone.localtime():%d/%m/%Y %H:%M}"])
    hoja.append([])
    for etiqueta, valor in _resumen_general(presupuesto):
        hoja.append([etiqueta, float(valor)])
        hoja.cell(row=hoja.max_row, column=1).font = Font(bold=True)

    hoja.append([])
    encabezado("Categoria", "Comprometido", "Ejecutado", "Total")
    for etiqueta, comprometido, ejecutado in _resumen_por_categoria(presupuesto):
        hoja.append([etiqueta, float(comprometido), float(ejecutado), float(comprometido + ejecutado)])

    hoja.append([])
    encabezado("Fecha", "Categoria", "Estado", "Descripcion", "Monto", "Registrado por", "Comprobante")
    for gasto in _gastos_ordenados(presupuesto):
        agregar_fila(hoja, [
            timezone.localtime(gasto.creado_en).strftime("%d/%m/%Y"),
            gasto.get_categoria_display(),
            gasto.get_estado_display(),
            gasto.descripcion,
            float(gasto.monto),
            gasto.registrado_por.username if gasto.registrado_por else "",
            "Si" if gasto.comprobante else "No",
        ])
    for columna, ancho in zip("ABCDEFG", (16, 18, 16, 40, 14, 18, 12)):
        hoja.column_dimensions[columna].width = ancho

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def exportar_pdf(presupuesto):
    """Reporte de gastos del reto en PDF con el encabezado de marca de los reportes (HU07)."""
    import io

    from django.utils import timezone
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from xml.sax.saxutils import escape

    from apps.reportes.services import VERDE_EAN, encabezado_de_marca

    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=22 * mm, bottomMargin=18 * mm,
        title=f"Presupuesto - {presupuesto.reto.titulo}", author="RetosEAN - Universidad EAN",
    )
    estilos = getSampleStyleSheet()
    estilo_tabla = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), VERDE_EAN),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ])

    def pesos(valor):
        return f"${valor:,.2f}"

    def tabla(filas, **kwargs):
        resultado = Table(filas, repeatRows=1, **kwargs)
        resultado.setStyle(estilo_tabla)
        return resultado

    detalle = [["Fecha", "Categoria", "Estado", "Descripcion", "Monto"]] + [
        [
            timezone.localtime(gasto.creado_en).strftime("%d/%m/%Y"),
            gasto.get_categoria_display(),
            gasto.get_estado_display(),
            Paragraph(escape(gasto.descripcion or "-"), estilos["Normal"]),
            pesos(gasto.monto),
        ]
        for gasto in _gastos_ordenados(presupuesto)
    ]
    resumen = _resumen_general(presupuesto)
    documento.build([
        Paragraph(escape(f"Presupuesto del reto: {presupuesto.reto.titulo}"), estilos["Title"]),
        Paragraph(f"Generado el {timezone.localtime():%d/%m/%Y %H:%M}", estilos["Normal"]),
        Spacer(1, 6 * mm),
        tabla([
            [etiqueta for etiqueta, _ in resumen],
            [pesos(v) if etiqueta != "% asignado" else f"{v:.1f}%" for etiqueta, v in resumen],
        ]),
        Spacer(1, 6 * mm),
        Paragraph("Totales por categoria", estilos["Heading3"]),
        tabla([["Categoria", "Comprometido", "Ejecutado", "Total"]] + [
            [etiqueta, pesos(comprometido), pesos(ejecutado), pesos(comprometido + ejecutado)]
            for etiqueta, comprometido, ejecutado in _resumen_por_categoria(presupuesto)
        ]),
        Spacer(1, 6 * mm),
        Paragraph("Detalle de gastos", estilos["Heading3"]),
        tabla(detalle, colWidths=[22 * mm, 28 * mm, 24 * mm, 76 * mm, 30 * mm]),
    ], onFirstPage=encabezado_de_marca, onLaterPages=encabezado_de_marca)
    return buffer.getvalue()
