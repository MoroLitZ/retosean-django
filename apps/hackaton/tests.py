from datetime import timedelta
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from apps.notificaciones.models import Notificacion
from apps.participaciones.models import Equipo, MiembroEquipo, Postulacion
from apps.retos.models import Reto
from apps.seguimiento.models import IntegracionAcademica
from apps.usuarios.models import Usuario

from . import services
from .models import EtapaHackaton, Hackathon, InscripcionHackaton, Jurado, VotacionHackaton


class BaseHackathonTests(TestCase):
    def setUp(self):
        self.admin = Usuario.objects.create_superuser(
            username="admin-hack", password="x", email="admin@ean.edu.co"
        )
        self.empresa = Usuario.objects.create_user(
            username="empresa-hack", password="x", rol="EMPRESA"
        )
        self.profesor = Usuario.objects.create_user(
            username="profe-hack", password="x", rol="PROFESOR"
        )
        self.reto = Reto.objects.create(
            empresa=self.empresa, titulo="Reto hackathon", estado="en_curso", tipo="hackathon"
        )
        IntegracionAcademica.objects.create(
            reto=self.reto, profesor=self.profesor, estado="aprobada"
        )
        self.hackathon = Hackathon.objects.create(reto=self.reto, estado="publicado")

    def _equipo(self, nombre, *estudiantes):
        equipo = Equipo.objects.create(reto=self.reto, nombre=nombre)
        for estudiante in estudiantes:
            MiembroEquipo.objects.create(equipo=equipo, estudiante=estudiante)
        return equipo

    def _estudiante(self, username):
        estudiante = Usuario.objects.create_user(username=username, password="x", rol="ESTUDIANTE")
        Postulacion.objects.create(reto=self.reto, estudiante=estudiante, estado="ACEPTADA")
        return estudiante

    def _jurado(self, username):
        usuario = Usuario.objects.create_user(username=username, password="x", rol="PROFESOR")
        return Jurado.objects.create(hackathon=self.hackathon, usuario=usuario, nombre=username)

    def _inscribir(self, equipo, estado="ACEPTADA"):
        return InscripcionHackaton.objects.create(
            hackathon=self.hackathon, equipo=equipo, estado=estado
        )


class EtapasTests(BaseHackathonTests):
    def test_crea_las_cuatro_etapas_del_backlog(self):
        services.crear_etapas_por_defecto(self.hackathon)
        tipos = list(self.hackathon.etapas.values_list("tipo", flat=True))
        self.assertEqual(
            sorted(tipos), ["desarrollo", "inscripcion", "premiacion", "presentacion"]
        )

    def test_no_duplica_etapas_si_ya_existen(self):
        services.crear_etapas_por_defecto(self.hackathon)
        services.crear_etapas_por_defecto(self.hackathon)
        self.assertEqual(self.hackathon.etapas.count(), 4)

    def test_etapa_fuera_de_fecha_no_esta_vigente(self):
        ayer = timezone.now() - timedelta(days=2)
        etapa = EtapaHackaton.objects.create(
            hackathon=self.hackathon, tipo="inscripcion", titulo="Inscripcion",
            inicia_en=ayer, termina_en=ayer + timedelta(days=1),
        )
        self.assertFalse(etapa.vigente)
        self.assertFalse(services.inscripciones_abiertas(self.hackathon))

    def test_etapa_en_curso_abre_las_inscripciones(self):
        ahora = timezone.now()
        EtapaHackaton.objects.create(
            hackathon=self.hackathon, tipo="inscripcion", titulo="Inscripcion",
            inicia_en=ahora - timedelta(days=1), termina_en=ahora + timedelta(days=5),
        )
        self.assertTrue(services.inscripciones_abiertas(self.hackathon))


class InscripcionTests(BaseHackathonTests):
    def test_estudiante_inscribe_su_equipo(self):
        estudiante = self._estudiante("est-insc")
        self._equipo("Equipo Alfa", estudiante)
        self.client.force_login(estudiante)
        equipo_id = Equipo.objects.get(nombre="Equipo Alfa").pk
        response = self.client.post(
            reverse("hackaton:inscribir_equipo", args=[self.hackathon.pk]),
            {"equipo": equipo_id},
        )
        self.assertEqual(response.status_code, 302)
        inscripcion = InscripcionHackaton.objects.get(equipo_id=equipo_id)
        self.assertEqual(inscripcion.estado, "PENDIENTE")

    def test_no_se_puede_inscribir_dos_veces_el_mismo_equipo(self):
        estudiante = self._estudiante("est-doble")
        equipo = self._equipo("Equipo Beta", estudiante)
        self._inscribir(equipo, estado="PENDIENTE")
        self.client.force_login(estudiante)
        response = self.client.post(
            reverse("hackaton:inscribir_equipo", args=[self.hackathon.pk]),
            {"equipo": equipo.pk},
        )
        # El equipo ya no aparece en el queryset del formulario.
        self.assertEqual(response.status_code, 200)
        self.assertEqual(InscripcionHackaton.objects.filter(equipo=equipo).count(), 1)

    def test_gestor_acepta_la_inscripcion(self):
        equipo = self._equipo("Equipo Gamma")
        inscripcion = self._inscribir(equipo, estado="PENDIENTE")
        self.client.force_login(self.admin)
        self.client.post(
            reverse("hackaton:gestionar_inscripcion", args=[inscripcion.pk]),
            {"accion": "aceptar"},
        )
        inscripcion.refresh_from_db()
        self.assertEqual(inscripcion.estado, "ACEPTADA")


class VotacionTests(BaseHackathonTests):
    def test_voto_fuera_de_rango_es_invalido(self):
        equipo = self._equipo("Equipo Delta")
        self._inscribir(equipo)
        jurado = self._jurado("jurado-uno")
        voto = VotacionHackaton(
            hackathon=self.hackathon, jurado=jurado, equipo=equipo, puntaje=Decimal("7.00")
        )
        with self.assertRaises(ValidationError):
            voto.full_clean()

    def test_no_se_puede_votar_un_equipo_sin_inscripcion_aceptada(self):
        equipo = self._equipo("Equipo Epsilon")
        self._inscribir(equipo, estado="PENDIENTE")
        jurado = self._jurado("jurado-dos")
        voto = VotacionHackaton(
            hackathon=self.hackathon, jurado=jurado, equipo=equipo, puntaje=Decimal("4.00")
        )
        with self.assertRaises(ValidationError):
            voto.full_clean()

    def test_jurado_registra_y_actualiza_su_voto(self):
        equipo = self._equipo("Equipo Zeta")
        self._inscribir(equipo)
        jurado = self._jurado("jurado-tres")
        self.client.force_login(jurado.usuario)
        url = reverse("hackaton:votar", args=[self.hackathon.pk])

        self.client.post(url, {"equipo": equipo.pk, "puntaje": "4.00", "comentario": "Bien"})
        self.assertEqual(VotacionHackaton.objects.count(), 1)

        self.client.post(url, {"equipo": equipo.pk, "puntaje": "4.75", "comentario": "Mejor"})
        self.assertEqual(VotacionHackaton.objects.count(), 1)
        self.assertEqual(VotacionHackaton.objects.get().puntaje, Decimal("4.75"))

    def test_quien_no_es_jurado_no_puede_votar(self):
        self.client.force_login(self.profesor)
        response = self.client.get(reverse("hackaton:votar", args=[self.hackathon.pk]))
        self.assertRedirects(response, reverse("hackaton:detalle", args=[self.hackathon.pk]))


class RankingTests(BaseHackathonTests):
    def setUp(self):
        super().setUp()
        self.equipo_a = self._equipo("Equipo A")
        self.equipo_b = self._equipo("Equipo B")
        self._inscribir(self.equipo_a)
        self._inscribir(self.equipo_b)
        self.jurado_1 = self._jurado("j1")
        self.jurado_2 = self._jurado("j2")

    def _votar(self, jurado, equipo, puntaje):
        VotacionHackaton.objects.create(
            hackathon=self.hackathon, jurado=jurado, equipo=equipo, puntaje=Decimal(puntaje)
        )

    def test_ordena_por_promedio_no_por_suma(self):
        # A: un solo voto de 5.0 (promedio 5). B: dos votos de 4.0 (promedio 4, suma 8).
        self._votar(self.jurado_1, self.equipo_a, "5.00")
        self._votar(self.jurado_1, self.equipo_b, "4.00")
        self._votar(self.jurado_2, self.equipo_b, "4.00")

        tabla = services.ranking(self.hackathon)
        self.assertEqual(tabla[0]["equipo"], "Equipo A")
        self.assertEqual(tabla[0]["posicion"], 1)
        self.assertEqual(tabla[1]["equipo"], "Equipo B")
        # Aunque B suma mas puntos totales.
        self.assertGreater(tabla[1]["total"], tabla[0]["total"])

    def test_desempate_por_numero_de_votos(self):
        self._votar(self.jurado_1, self.equipo_a, "4.00")
        self._votar(self.jurado_1, self.equipo_b, "4.00")
        self._votar(self.jurado_2, self.equipo_b, "4.00")
        tabla = services.ranking(self.hackathon)
        self.assertEqual(tabla[0]["equipo"], "Equipo B")

    def test_equipos_sin_votar_del_jurado(self):
        self._votar(self.jurado_1, self.equipo_a, "4.00")
        pendientes = services.equipos_sin_votar(self.hackathon, self.jurado_1)
        self.assertEqual([i.equipo_id for i in pendientes], [self.equipo_b.pk])

    def test_resultados_ocultos_hasta_publicarlos(self):
        estudiante = self._estudiante("est-ranking")
        self.client.force_login(estudiante)
        response = self.client.get(reverse("hackaton:ranking", args=[self.hackathon.pk]))
        self.assertFalse(response.context["puede_ver"])

    def test_publicar_resultados_los_hace_visibles_y_notifica(self):
        estudiante = self._estudiante("est-ganador")
        MiembroEquipo.objects.create(equipo=self.equipo_a, estudiante=estudiante)
        self._votar(self.jurado_1, self.equipo_a, "5.00")

        services.publicar_resultados(self.hackathon, self.admin)
        self.hackathon.refresh_from_db()
        self.assertTrue(self.hackathon.resultados_publicados)
        self.assertTrue(
            Notificacion.objects.filter(
                usuario=estudiante, evento="HACKATON_RESULTADOS"
            ).exists()
        )

        self.client.force_login(estudiante)
        response = self.client.get(reverse("hackaton:ranking", args=[self.hackathon.pk]))
        self.assertTrue(response.context["puede_ver"])


class PublicacionTests(BaseHackathonTests):
    def test_publicar_notifica_a_los_estudiantes_aceptados(self):
        estudiante = self._estudiante("est-aviso")
        self.hackathon.estado = "borrador"
        self.hackathon.save()

        services.publicar(self.hackathon, self.admin)
        self.hackathon.refresh_from_db()
        self.assertEqual(self.hackathon.estado, "publicado")
        self.assertTrue(
            Notificacion.objects.filter(
                usuario=estudiante, evento="HACKATON_PUBLICADO"
            ).exists()
        )

    def test_profesor_ajeno_no_gestiona_el_hackathon(self):
        ajeno = Usuario.objects.create_user(username="profe-ajeno", password="x", rol="PROFESOR")
        self.client.force_login(ajeno)
        response = self.client.post(reverse("hackaton:publicar", args=[self.hackathon.pk]))
        self.assertRedirects(response, reverse("hackaton:detalle", args=[self.hackathon.pk]))
        self.hackathon.refresh_from_db()
        self.assertEqual(self.hackathon.estado, "publicado")


class CalendarioAutomaticoHU09Tests(BaseHackathonTests):
    """HU09: avance por calendario y publicacion automatica de resultados al cierre."""

    def _etapa(self, tipo, inicia, termina=None):
        return EtapaHackaton.objects.create(
            hackathon=self.hackathon, tipo=tipo, titulo=tipo, inicia_en=inicia, termina_en=termina
        )

    def test_pasa_a_en_curso_al_iniciar_el_desarrollo(self):
        ahora = timezone.now()
        self._etapa("desarrollo", ahora - timedelta(hours=1))
        self._etapa("premiacion", ahora + timedelta(days=5))
        self.assertEqual(services.avanzar_por_calendario(self.hackathon), ["en_curso"])
        self.hackathon.refresh_from_db()
        self.assertEqual(self.hackathon.estado, "en_curso")
        self.assertFalse(self.hackathon.resultados_publicados)

    def test_al_cierre_publica_resultados_y_ganador_y_deja_historial(self):
        from apps.seguimiento.services import historial_de

        ahora = timezone.now()
        self._etapa("desarrollo", ahora - timedelta(days=3))
        self._etapa("premiacion", ahora - timedelta(minutes=5))
        estudiante = self._estudiante("est-cierre")
        equipo = self._equipo("Equipo Cierre", estudiante)
        self._inscribir(equipo)
        VotacionHackaton.objects.create(hackathon=self.hackathon, jurado=self._jurado("jur-cierre"),
                                        equipo=equipo, puntaje=Decimal("4.5"))

        services.avanzar_por_calendario(self.hackathon)
        self.hackathon.refresh_from_db()
        self.assertEqual(self.hackathon.estado, "finalizado")
        self.assertTrue(self.hackathon.resultados_publicados)
        aviso = Notificacion.objects.get(usuario=estudiante, evento="HACKATON_RESULTADOS")
        self.assertIn("Equipo Cierre", aviso.mensaje)
        pasos = list(historial_de(self.hackathon).values_list("estado_nuevo", flat=True))
        self.assertIn("finalizado", pasos)
        # Idempotente: una segunda pasada no repite nada.
        self.assertEqual(services.avanzar_por_calendario(self.hackathon), [])

    def test_sin_fechas_no_cierra_solo(self):
        self._etapa("premiacion", None)
        self.assertEqual(services.avanzar_por_calendario(self.hackathon), [])

    def test_sin_premiacion_cierra_al_terminar_la_ultima_etapa(self):
        ahora = timezone.now()
        self._etapa("presentacion", ahora - timedelta(days=2), ahora - timedelta(hours=1))
        self.assertEqual(services.fecha_de_cierre(self.hackathon), ahora - timedelta(hours=1))
        services.avanzar_por_calendario(self.hackathon)
        self.hackathon.refresh_from_db()
        self.assertTrue(self.hackathon.resultados_publicados)

    def test_el_comando_programado_avanza_los_hackathones(self):
        from io import StringIO

        from django.core.management import call_command

        self._etapa("premiacion", timezone.now() - timedelta(minutes=1))
        salida = StringIO()
        call_command("avanzar_hackathones", stdout=salida)
        self.hackathon.refresh_from_db()
        self.assertTrue(self.hackathon.resultados_publicados)
        self.assertIn("1 hackathon(es) actualizados", salida.getvalue())

    def test_abrir_el_detalle_aplica_el_calendario(self):
        self._etapa("premiacion", timezone.now() - timedelta(minutes=1))
        self.client.force_login(self.admin)
        self.client.get(reverse("hackaton:detalle", args=[self.hackathon.pk]))
        self.hackathon.refresh_from_db()
        self.assertTrue(self.hackathon.resultados_publicados)


class LogisticaEInformeHU09Tests(BaseHackathonTests):
    def test_gestor_registra_logistica_e_invitados(self):
        self.client.force_login(self.admin)
        respuesta = self.client.post(reverse("hackaton:logistica", args=[self.hackathon.pk]), {
            "modalidad": "hibrida", "lugar": "Auditorio EAN", "enlace_transmision": "https://youtube.com/live/x",
            "notas_logistica": "Refrigerios a las 10",
            "invitados-TOTAL_FORMS": "1", "invitados-INITIAL_FORMS": "0",
            "invitados-MIN_NUM_FORMS": "0", "invitados-MAX_NUM_FORMS": "1000",
            "invitados-0-tipo": "MENTOR", "invitados-0-nombre": "Ana Mentora",
            "invitados-0-organizacion": "ACME", "invitados-0-email": "ana@acme.co", "invitados-0-confirmado": "on",
        })
        self.assertRedirects(respuesta, reverse("hackaton:detalle", args=[self.hackathon.pk]))
        self.hackathon.refresh_from_db()
        self.assertEqual(self.hackathon.lugar, "Auditorio EAN")
        invitado = self.hackathon.invitados.get()
        self.assertTrue(invitado.confirmado)
        detalle = self.client.get(reverse("hackaton:detalle", args=[self.hackathon.pk]))
        self.assertContains(detalle, "Ana Mentora")

    def test_evento_virtual_exige_enlace_de_transmision(self):
        from .forms import LogisticaForm

        form = LogisticaForm({"modalidad": "virtual"}, instance=self.hackathon)
        self.assertFalse(form.is_valid())
        self.assertIn("enlace_transmision", form.errors)

    def test_estudiante_no_gestiona_la_logistica(self):
        estudiante = self._estudiante("est-log")
        self.client.force_login(estudiante)
        respuesta = self.client.get(reverse("hackaton:logistica", args=[self.hackathon.pk]))
        self.assertEqual(respuesta.status_code, 302)

    def test_informe_final_en_pdf_y_excel_con_el_ranking(self):
        from io import BytesIO

        from openpyxl import load_workbook

        estudiante = self._estudiante("est-inf")
        equipo = self._equipo("Equipo =SUMA(A1)", estudiante)
        self._inscribir(equipo)
        VotacionHackaton.objects.create(hackathon=self.hackathon, jurado=self._jurado("jur-inf"),
                                        equipo=equipo, puntaje=Decimal("5"))
        services.publicar_resultados(self.hackathon, self.admin)
        url = reverse("hackaton:informe_final", args=[self.hackathon.pk])

        # La empresa duena del reto tambien puede descargarlo.
        self.client.force_login(self.empresa)
        pdf = self.client.get(url + "?formato=pdf")
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))

        excel = self.client.get(url + "?formato=excel")
        libro = load_workbook(BytesIO(excel.content))
        self.assertIn("Ranking", libro.sheetnames)
        valores = [c.value for fila in libro["Ranking"].iter_rows() for c in fila]
        self.assertIn("Equipo =SUMA(A1)", valores)
        self.assertEqual(libro["Ranking"].cell(row=5, column=2).data_type, "s")

    def test_estudiante_no_descarga_el_informe(self):
        estudiante = self._estudiante("est-noinf")
        self.client.force_login(estudiante)
        respuesta = self.client.get(reverse("hackaton:informe_final", args=[self.hackathon.pk]))
        self.assertEqual(respuesta.status_code, 302)
