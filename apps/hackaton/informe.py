"""Informe final exportable del hackathon (HU09), en PDF y Excel.

Reusa la marca y las protecciones de apps.reportes (encabezado EAN, celdas
que no se convierten en formulas) para que se vea como el resto de reportes.
"""

import io
from xml.sax.saxutils import escape

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.reportes.services import GRIS_SUAVE, GRIS_TEXTO, VERDE_EAN, agregar_fila, encabezado_de_marca

from .services import ranking


def _fecha(valor, con_hora=True):
    if not valor:
        return "-"
    if hasattr(valor, "tzinfo"):
        valor = timezone.localtime(valor)
    return valor.strftime("%d/%m/%Y %H:%M" if con_hora else "%d/%m/%Y")


def datos_informe(hackathon):
    """Secciones del informe como (titulo, encabezados, filas); igual para PDF y Excel."""
    reto = hackathon.reto
    tabla = ranking(hackathon)
    ganador = tabla[0]["equipo"] if tabla and hackathon.resultados_publicados else None

    convocatoria = [
        ("Reto", reto.titulo),
        ("Codigo", reto.codigo or "-"),
        ("Empresa", reto.empresa.get_full_name() or reto.empresa.username),
        ("Estado", hackathon.get_estado_display()),
        ("Fechas", f"{_fecha(reto.fecha_inicio_tentativa, False)} a {_fecha(reto.fecha_fin_tentativa, False)}"),
        ("Descripcion", reto.descripcion or "-"),
        ("Reglas", hackathon.reglas or "-"),
        ("Premios", reto.premios or "-"),
        ("Resultados publicados", "Si" if hackathon.resultados_publicados else "No"),
        ("Equipo ganador", ganador or "-"),
    ]
    logistica = [
        ("Modalidad", hackathon.get_modalidad_display()),
        ("Espacio fisico", hackathon.lugar or "-"),
        ("Transmision", hackathon.enlace_transmision or "-"),
        ("Notas", hackathon.notas_logistica or "-"),
    ]
    inscripciones = hackathon.inscripciones.select_related("equipo").prefetch_related("equipo__miembros__estudiante")

    return [
        ("Convocatoria", ["Campo", "Valor"], convocatoria),
        ("Logistica del evento", ["Campo", "Valor"], logistica),
        ("Etapas", ["Etapa", "Tipo", "Inicia", "Termina"], [
            (e.titulo, e.get_tipo_display(), _fecha(e.inicia_en), _fecha(e.termina_en))
            for e in hackathon.etapas.all()
        ]),
        ("Invitados", ["Nombre", "Tipo", "Organizacion", "Confirmado"], [
            (i.nombre, i.get_tipo_display(), i.organizacion or "-", "Si" if i.confirmado else "No")
            for i in hackathon.invitados.all()
        ]),
        ("Jurados", ["Nombre", "Correo", "Especialidad"], [
            (j.nombre, j.email or "-", j.especialidad or "-") for j in hackathon.jurados.all()
        ]),
        ("Equipos inscritos", ["Equipo", "Estado", "Integrantes", "Inscrito"], [
            (
                i.equipo.nombre, i.get_estado_display(),
                ", ".join(m.estudiante.get_full_name() or m.estudiante.username for m in i.equipo.miembros.all()) or "-",
                _fecha(i.creado_en),
            )
            for i in inscripciones
        ]),
        ("Ranking", ["Puesto", "Equipo", "Promedio", "Total", "Votos"], [
            (f["posicion"], f["equipo"], f"{f['promedio']:.2f}", f"{f['total']:.2f}", f["votos"]) for f in tabla
        ]),
    ]


def informe_pdf(hackathon):
    buffer = io.BytesIO()
    titulo = f"Informe final del hackathon: {hackathon.reto.titulo}"
    documento = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=22 * mm, bottomMargin=18 * mm, title=titulo, author="RetosEAN - Universidad EAN",
    )
    estilos = getSampleStyleSheet()
    celda = ParagraphStyle("Celda", parent=estilos["Normal"], fontSize=8, leading=10)
    cabecera = ParagraphStyle("Cabecera", parent=celda, textColor=colors.white, fontName="Helvetica-Bold")
    elementos = [
        Paragraph(escape(titulo), ParagraphStyle("T", parent=estilos["Title"], fontSize=15, alignment=0)),
        Paragraph(f"Generado el {timezone.localtime():%d/%m/%Y %H:%M}",
                  ParagraphStyle("M", parent=estilos["Normal"], fontSize=8.5, textColor=GRIS_TEXTO)),
        Spacer(1, 8),
    ]
    for seccion, encabezados, filas in datos_informe(hackathon):
        elementos.append(Paragraph(seccion, estilos["Heading3"]))
        if not filas:
            elementos.append(Paragraph("Sin registros.", celda))
            continue
        datos = [[Paragraph(e, cabecera) for e in encabezados]]
        datos += [[Paragraph(escape(str(v)), celda) for v in fila] for fila in filas]
        tabla = Table(datos, repeatRows=1, colWidths=[documento.width / len(encabezados)] * len(encabezados)
                      if len(encabezados) > 2 else [documento.width * 0.28, documento.width * 0.72])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), VERDE_EAN),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d5db")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, GRIS_SUAVE]),
        ]))
        elementos += [tabla, Spacer(1, 6)]
    documento.build(elementos, onFirstPage=encabezado_de_marca, onLaterPages=encabezado_de_marca)
    return buffer.getvalue()


def informe_excel(hackathon):
    libro = Workbook()
    libro.remove(libro.active)
    for seccion, encabezados, filas in datos_informe(hackathon):
        hoja = libro.create_sheet(seccion[:31])
        hoja.append(["retosean · Universidad EAN"])
        hoja["A1"].font = Font(bold=True, color="3BAC52")
        agregar_fila(hoja, [f"{seccion} - {hackathon.reto.titulo}"])
        hoja["A2"].font = Font(bold=True, size=13)
        hoja.append([])
        hoja.append(encabezados)
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="3BAC52")
        for fila in filas:
            agregar_fila(hoja, list(fila))
        for columna in "ABCDE":
            hoja.column_dimensions[columna].width = 28
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
